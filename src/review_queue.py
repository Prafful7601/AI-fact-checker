"""Appends abstained claims to results/review_queue.xlsx (created on first use).
Columns match the spec exactly: claim_id, claim, table_id, verdict,
agreement_score, executed_checks, reason, status, created_at."""
from __future__ import annotations
import json
import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUMNS = ["claim_id", "claim", "table_id", "verdict", "agreement_score",
           "executed_checks", "reason", "status", "created_at"]

_lock = threading.Lock()


def _ensure_workbook(path: Path):
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "review_queue"
    ws.append(COLUMNS)
    wb.save(path)


def append_review_row(path: str, claim_id: str, claim: str, table_id: str,
                       verdict: str, agreement_score: float, executed_checks: list,
                       reason: str) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with _lock:
        _ensure_workbook(p)
        wb = load_workbook(p)
        ws = wb["review_queue"]
        ws.append([
            claim_id, claim, table_id, verdict, agreement_score,
            json.dumps(executed_checks, default=str)[:32000],  # Excel cell limit
            reason, "Pending", datetime.now(timezone.utc).isoformat(),
        ])
        wb.save(p)
